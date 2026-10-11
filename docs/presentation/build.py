"""Build an editable Russian project deck using native PowerPoint shapes."""

import argparse
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
import qrcode


ROOT = Path(__file__).resolve().parents[2]
ASSETS = Path(__file__).resolve().parent / "assets"
FONTS = ROOT / "apps/web/static/fonts"
WIDTH, HEIGHT = 13.333333, 7.5
TOTAL = 16
PURPLE = "4F0B78"
DARK = "2D1044"
INK = "25232B"
MUTED = "776C83"
BG = "F7F4FA"
PALE = "EFE7F5"
ROSE = "FFD4E7"
WHITE = "FFFFFF"
GREEN = "287A67"
DEMO = "https://brokenbrain.online"
REPO = "https://github.com/06ED/fsp-career"
YANDEX = "https://disk.yandex.ru/d/8Kw-km6dKgXJEQ"
GOOGLE = "https://drive.google.com/drive/folders/1nd1SiafMHe-dXMK5UbNHw3jTvEp_fZPU?usp=sharing"


def rgb(value):
    return RGBColor.from_string(value)


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(WIDTH)
        self.prs.slide_height = Inches(HEIGHT)
        self.prs.core_properties.title = "FSP Career — презентация проекта"
        self.prs.core_properties.subject = "Оценка навыков и объяснимый подбор специалистов"
        self.prs.core_properties.author = "FSP Career"
        self.prs.core_properties.language = "ru-RU"
        self.prs.core_properties.keywords = "FSP Career, CAT, IRT, Ktor, SvelteKit, GitOps"
        self.prs.core_properties.comments = "Редактируемые объекты; исходник в docs/presentation/build.py."
        self.text_boxes = []
        self.fit_errors = []

    def rect(self, slide, x, y, w, h, fill=WHITE, stroke=None, rounded=True):
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        if rounded:
            shape.adjustments[0] = 0.12
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if stroke:
            shape.line.color.rgb = rgb(stroke)
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    def line(self, slide, x1, y1, x2, y2, color=PURPLE, width=2):
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
        )
        line.line.color.rgb = rgb(color)
        line.line.width = Pt(width)
        return line

    def wrap(self, text, width_pt, size, family, bold):
        path = FONTS / f"{family}-{'Bold' if bold else 'Regular'}.ttf"
        font = ImageFont.truetype(str(path), round(size * 4))
        lines = []
        for paragraph in text.split("\n"):
            words = paragraph.split()
            if not words:
                lines.append("")
                continue
            current = words[0]
            for word in words[1:]:
                trial = f"{current} {word}"
                if font.getlength(trial) / 4 <= width_pt:
                    current = trial
                else:
                    lines.append(current)
                    current = word
            lines.append(current)
        if any(font.getlength(line) / 4 > width_pt + 1 for line in lines):
            return None
        return lines

    def text(self, slide, text, x, y, w, h, size=21, color=INK, bold=False,
             family="Montserrat", align=PP_ALIGN.LEFT, url=None, minimum=16):
        requested = size
        while size >= minimum:
            lines = self.wrap(text, w * 72 - 2, size, family, bold)
            if lines is not None and len(lines) * size * 1.16 <= h * 72 - 2:
                break
            size -= 0.5
        else:
            self.fit_errors.append(f"Slide {len(self.prs.slides)}: {text!r}, box={w}×{h}, max={requested}")
            size = minimum
            lines = self.wrap(text, w * 72 - 2, size, family, bold) or text.split("\n")
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = False
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.TOP
        for index, value in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = align
            paragraph.space_before = paragraph.space_after = Pt(0)
            paragraph.line_spacing = 1.05
            run = paragraph.add_run()
            run.text = value
            run.font.name = family
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
        if url:
            # Shape-level links preserve the explicit text color in Impress/PDF.
            shape.click_action.hyperlink.address = url
        self.text_boxes.append((len(self.prs.slides), text, size, len(lines)))
        return shape

    def pill(self, slide, label, x, y, w, fill=PALE, color=PURPLE):
        self.rect(slide, x, y, w, 0.4, fill)
        self.text(slide, label, x + 0.1, y + 0.08, w - 0.2, 0.25,
                  size=12, minimum=11, bold=True, color=color, align=PP_ALIGN.CENTER)

    def slide(self, title, subtitle, notes, dark=False, section="ПРОЕКТ"):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(DARK if dark else BG)
        n = len(self.prs.slides)
        color = WHITE if dark else PURPLE
        muted = "DAC5E6" if dark else MUTED
        self.text(slide, f"ФСП / КАРЬЕРА  ·  {section}", 0.55, 0.28, 11.4, 0.3,
                  size=11.5, minimum=11.5, color=ROSE if dark else MUTED, bold=True)
        if title:
            self.text(slide, title, 0.55, 0.78, 12.2, 0.65, size=34,
                      family="Play", bold=True, color=color, minimum=27)
        if subtitle:
            self.text(slide, subtitle, 0.57, 1.46, 12.1, 0.65,
                      size=17.5, minimum=16, color=muted)
        self.line(slide, 0.55, 7.05, 12.78, 7.05, "67457D" if dark else "E6DDEC", 0.7)
        self.text(slide, "FSP Career · brokenbrain.online", 0.55, 7.15, 10.5, 0.2,
                  size=10, minimum=10, color=muted, url=DEMO)
        self.text(slide, f"{n:02d} / {TOTAL}", 11.6, 7.13, 1.18, 0.25,
                  size=11.5, minimum=11.5, color=muted, align=PP_ALIGN.RIGHT)
        slide.notes_slide.notes_text_frame.text = notes
        return slide

    def card(self, slide, x, y, w, h, label, title, body, fill=WHITE):
        self.rect(slide, x, y, w, h, fill)
        self.text(slide, label, x + 0.25, y + 0.2, w - 0.5, 0.35,
                  size=12, minimum=11, color=PURPLE, bold=True)
        self.text(slide, title, x + 0.25, y + 0.69, w - 0.5, 0.88,
                  size=25, minimum=20, family="Play", bold=True, color=PURPLE)
        self.text(slide, body, x + 0.25, y + 1.76, w - 0.5, h - 1.98,
                  size=20, minimum=17.5)

    def flow(self, slide, labels, y, colors=None):
        gap, x, total_w = 0.29, 0.55, 12.23
        w = (total_w - gap * (len(labels) - 1)) / len(labels)
        for index, (title, body) in enumerate(labels):
            fill = (colors or [WHITE] * len(labels))[index]
            self.rect(slide, x, y, w, 1.32, fill)
            self.text(slide, title, x + 0.13, y + 0.18, w - 0.26, 0.38,
                      size=20, minimum=17, family="Play", bold=True, color=PURPLE)
            self.text(slide, body, x + 0.13, y + 0.68, w - 0.26, 0.48,
                      size=15.5, minimum=13.5, color=MUTED)
            if index < len(labels) - 1:
                self.text(slide, "→", x + w + 0.04, y + 0.46, gap - 0.06, 0.35,
                          size=17, minimum=16, color=PURPLE, align=PP_ALIGN.CENTER)
            x += w + gap

    def image(self, slide, path, x, y, w, h, crop=None):
        with Image.open(path) as original:
            image = original.convert("RGB")
            if crop:
                image = image.crop(crop)
            iw, ih = image.size
            ratio = min(w / iw, h / ih)
            sw, sh = iw * ratio, ih * ratio
            stream = BytesIO()
            image.save(stream, format="PNG")
            stream.seek(0)
            slide.shapes.add_picture(stream, Inches(x + (w - sw) / 2),
                                     Inches(y + (h - sh) / 2), Inches(sw), Inches(sh))

    def finish(self, output):
        if self.fit_errors:
            raise ValueError("Text does not fit:\n" + "\n".join(self.fit_errors))
        if len(self.prs.slides) != TOTAL:
            raise ValueError(f"Expected {TOTAL} slides, got {len(self.prs.slides)}")
        for n, slide in enumerate(self.prs.slides, 1):
            for shape in slide.shapes:
                if shape.left < 0 or shape.top < 0:
                    raise ValueError(f"Negative position on slide {n}")
                if shape.left + shape.width > self.prs.slide_width + 100:
                    raise ValueError(f"Horizontal overflow on slide {n}")
                if shape.top + shape.height > self.prs.slide_height + 100:
                    raise ValueError(f"Vertical overflow on slide {n}")
            if not slide.notes_slide.notes_text_frame.text.strip():
                raise ValueError(f"Speaker notes missing on slide {n}")
        self.prs.save(output)
        print(f"Created {TOTAL} slides, {len(self.text_boxes)} fitted text blocks: {output}")


def build(output):
    d = Deck()

    # 01 — Cover.
    s = d.slide("", "", "Представьте FSP Career: платформа связывает проверку навыков и поиск команды. "
                "В презентации показаны реализованные сценарии и техническая основа. Около 8–10 минут.",
                dark=True, section="ПРЕЗЕНТАЦИЯ ПРОЕКТА")
    d.text(s, "FSP Career", 0.65, 1.32, 8.0, 1.12, size=65, minimum=60,
           family="Play", bold=True, color=WHITE)
    d.text(s, "Навыки → возможности", 0.68, 2.66, 8.1, 0.7,
           size=34, minimum=30, color=ROSE, family="Play")
    d.text(s, "Оценка навыков и объяснимый подбор\nспециалистов под задачи команды", 0.7, 3.64,
           7.7, 1.25, size=26, minimum=23, color=WHITE)
    d.pill(s, "СПЕЦИАЛИСТ", 0.7, 5.37, 2.4, fill="67457D", color=WHITE)
    d.pill(s, "РАБОТОДАТЕЛЬ", 3.32, 5.37, 2.75, fill="67457D", color=WHITE)
    for y, number, title, body in [(1.42, "01", "Профиль", "Опыт и видимость"),
                                   (3.03, "02", "Оценка", "Категория по тесту"),
                                   (4.64, "03", "Подбор", "Причины совпадения")]:
        d.rect(s, 9.14, y, 3.63, 1.35, "4E2769")
        d.text(s, number, 9.4, y + 0.18, 0.5, 0.36, size=17, color=ROSE)
        d.text(s, title, 10.06, y + 0.16, 2.42, 0.45, size=26, family="Play", bold=True, color=WHITE)
        d.text(s, body, 9.41, y + 0.79, 3.05, 0.42, size=16, color=WHITE, minimum=14)

    # 02 — Problem.
    s = d.slide("Задача: связать навыки с потребностью команды", "Три вопроса, на которые отвечает платформа.",
                "Проблема состоит не в количестве резюме, а в сопоставлении заявленного опыта и задач команды. "
                "Не приводите неподтверждённые рыночные цифры: здесь качественная постановка задачи.")
    d.card(s, 0.55, 2.25, 3.89, 3.52, "СПЕЦИАЛИСТ", "Как показать уровень?",
           "Опыт и технологии в резюме требуют понятного подтверждения навыков.")
    d.card(s, 4.72, 2.25, 3.89, 3.52, "КОМПАНИЯ", "Кто подходит задачам?",
           "Нужны категория, релевантный опыт и объяснение позиции в подборке.")
    d.card(s, 8.89, 2.25, 3.89, 3.52, "ВЗАИМОДЕЙСТВИЕ", "Когда открывать контакт?",
           "Условия предложения и решение специалиста должны быть явными.")
    d.text(s, "Единый путь: описание опыта → проверка → сопоставление → предложение", 0.6, 6.13,
           12.1, 0.58, size=22, minimum=19, family="Play", color=PURPLE)

    # 03 — Solution.
    s = d.slide("Два кабинета — один сквозной сценарий", "Категория = специализация + подтверждённый грейд.",
                "Пройдите оба пути слева направо. Резюме описывает опыт, тест подтверждает категорию, "
                "работодатель задаёт потребность и отправляет предложение. Самооценка не присваивает грейд.")
    d.pill(s, "ПУТЬ СПЕЦИАЛИСТА", 0.55, 2.25, 3.35)
    d.flow(s, [("Профиль", "Опыт и согласия"), ("Опрос", "Направление"),
               ("CAT-тест", "Проверка навыков"), ("Категория", "Грейд по тесту"),
               ("Предложения", "Решение кандидата")], 2.96)
    d.pill(s, "ПУТЬ РАБОТОДАТЕЛЯ", 0.55, 4.68, 3.85, fill=ROSE)
    d.flow(s, [("Компания", "Контакты и задачи"), ("Потребность", "Грейд и стек"),
               ("Подбор", "Рейтинг и факторы"), ("Приглашение", "Условия и вилка"),
               ("Контакт", "После принятия")], 5.28)

    # 04 — Product screenshot.
    s = d.slide("Работающее приложение", "Публичная демонстрация: brokenbrain.online",
                "Это настоящий скриншот публичного лендинга, снятый без авторизованной сессии. "
                "Откройте демо либо скринкасты. Не выдавайте схемы следующих слайдов за скриншоты кабинетов.",
                section="ПРОДУКТ")
    d.rect(s, 0.55, 2.22, 8.58, 4.62)
    d.image(s, ASSETS / "landing.png", 0.67, 2.36, 8.34, 4.35)
    d.card(s, 9.38, 2.22, 3.4, 4.62, "ВЕБ-ПЛАТФОРМА", "От первого входа до предложения",
           "Две роли\n\nПрофиль и оценка\n\nПодбор и приглашения\n\nВакансии и задания", fill=PALE)

    # 05 — Candidate.
    s = d.slide("Специалист управляет своим профилем", "Опыт, документы и видимость — в одном кабинете.",
                "Покажите профиль, импорт PDF и приватность. PDF создаёт проверяемый черновик: пользователь "
                "исправляет поля и подтверждает. Контактный email редактируется отдельно от логина. "
                "Скрытые разделы исключаются из выдачи и факторов.", section="ПРОДУКТ")
    d.card(s, 0.55, 2.28, 3.89, 3.75, "ПРОФИЛЬ", "Опыт и технологии",
           "Специализация и отрасль\n\nИстория работы\n\nСтек и желаемые роли")
    d.card(s, 4.72, 2.28, 3.89, 3.75, "PDF-РЕЗЮМЕ", "Проверяемый импорт",
           "PDF → черновик\n\nРедактирование → подтверждение\n\nЭкспорт из профиля")
    d.card(s, 8.89, 2.28, 3.89, 3.75, "ПРИВАТНОСТЬ", "Явное управление",
           "Обработка данных\n\nПубликация профиля\n\nВидимость разделов")
    d.text(s, "Самооценка в профиле и подтверждённая категория хранятся отдельно", 0.65, 6.31,
           12.0, 0.45, size=21, minimum=19, color=PURPLE, family="Play")

    # 06 — Assessment.
    s = d.slide("Адаптивная оценка навыков", "Движок cat-v1: последовательные задания и единая шкала результата.",
                "12 заданий на 30 минут. CAT выбирает следующее семейство по информации и квотам тем. "
                "IRT/EAP оценивает уровень; итоговый балл — нормированный процентиль, а не доля верных ответов. "
                "Пересдачи ограничены 90 днями по специализации: после неудачи возможен уровень ниже, "
                "после уверенного успеха — выше, с проверкой истории. Новая неудача не понижает грейд.",
                section="АЛГОРИТМЫ")
    for x, number, label in [(0.55, "12", "заданий"), (4.72, "30", "минут"), (8.89, "5", "специализаций")]:
        d.rect(s, x, 2.23, 3.89, 1.61)
        d.text(s, number, x + 0.22, 2.42, 1.6, 0.85, size=50, family="Play", bold=True, color=PURPLE)
        d.text(s, label, x + 0.25, 3.3, 3.39, 0.34, size=17, minimum=15, color=MUTED)
    d.rect(s, 0.55, 4.13, 7.3, 2.52)
    d.text(s, "CAT / IRT", 0.82, 4.39, 6.6, 0.5, size=26, family="Play", bold=True, color=PURPLE)
    d.text(s, "Опрос задаёт направление и грейд.\nCAT подбирает задания по ответам.\nIRT/EAP рассчитывает итоговый уровень.",
           0.82, 5.12, 6.69, 1.32, size=20, minimum=18.5)
    d.rect(s, 8.14, 4.13, 4.64, 2.52, fill=PALE)
    d.text(s, "Junior · Middle · Senior", 8.41, 4.42, 4.08, 0.56,
           size=23, family="Play", bold=True, color=PURPLE)
    d.text(s, "Успех подтверждает грейд.\nПодтверждение сохраняется.\nПересдачи: 90 дней + условия.",
           8.41, 5.17, 4.08, 1.2, size=18.5, minimum=17)

    # 07 — Matching.
    s = d.slide("Подбор с объяснением каждого фактора", "Сначала категория, затем соответствие потребности команды.",
                "Это базовые веса алгоритма, не результат конкретного кандидата. Неприменимые факторы "
                "удаляются, веса нормируются. Итог умножается на соответствие грейда. "
                "В банке порядок по силе профиля: 65 тест + 25 ФСП + 10 актуальность. "
                "Самоописание влияет на релевантность, но не присваивает категорию.", section="АЛГОРИТМЫ")
    d.rect(s, 0.55, 2.22, 7.35, 4.48)
    for i, (label, weight) in enumerate([("Тест", 30), ("Текст опыта", 25), ("Стек", 20),
                                        ("Достижения ФСП", 15), ("Отрасль", 5), ("Актуальность", 5)]):
        y = 2.55 + i * 0.63
        d.text(s, label, 0.85, y, 2.8, 0.37, size=17.5, minimum=16)
        d.rect(s, 3.82, y + 0.07, 3.0, 0.22, fill=PALE)
        d.rect(s, 3.82, y + 0.07, weight / 30 * 3.0, 0.22, fill=PURPLE)
        d.text(s, str(weight), 7.02, y, 0.54, 0.37, size=18, bold=True, color=PURPLE)
    d.rect(s, 8.2, 2.22, 4.58, 4.48, fill=PALE)
    d.text(s, "Коэффициент грейда", 8.5, 2.56, 4.0, 0.7,
           size=25, minimum=21, family="Play", bold=True, color=PURPLE)
    d.text(s, "Точный уровень       × 1,00\nВыше на ступень      × 0,85\nВыше на две          × 0,70\nНиже, по запросу     × 0,60", 8.5, 3.58,
           4.0, 2.08, size=19, minimum=17.5)
    d.text(s, "В выдаче: рейтинг, объяснение,\nфакторы и карточка кандидата", 8.5, 5.88,
           4.0, 0.65, size=16.5, minimum=15, color=MUTED)

    # 08 — Privacy.
    s = d.slide("Контакты открываются по действию специалиста", "В сценарии приглашений доступ появляется после принятия предложения.",
                "Схема показывает правило доступа, а не данные реального пользователя. Работодатель "
                "видит категорию и разрешённые разделы до принятия, контакты после принятия — только своей компании. "
                "Самостоятельный отклик является отдельным основанием раскрытия этой компании.", section="ПРОДУКТ")
    for x, label, fill, color, email, phone in [(0.55, "ПРИГЛАШЕНИЕ ОТПРАВЛЕНО", WHITE, MUTED,
                                               "Email: скрыт", "Телефон: скрыт"),
                                              (7.14, "КАНДИДАТ ПРИНЯЛ", PALE, GREEN,
                                               "Email: contact@example.org", "Телефон: доступен компании")]:
        d.rect(s, x, 2.3, 5.64, 3.58, fill)
        d.text(s, label, x + 0.3, 2.58, 5.04, 0.34, size=12, minimum=11.5, bold=True, color=PURPLE)
        d.text(s, "Подтверждённая категория", x + 0.3, 3.16, 5.04, 0.67,
               size=25, minimum=21, family="Play", bold=True, color=PURPLE)
        d.text(s, f"{email}\n\n{phone}", x + 0.3, 4.18, 5.04, 1.27,
               size=21, minimum=18.5, color=color)
    d.text(s, "→", 6.23, 3.68, 0.73, 0.7, size=36, color=PURPLE, align=PP_ALIGN.CENTER)
    d.text(s, "Скрытые разделы исключаются из рейтинга. Самоотклик раскрывает контакт автору вакансии.",
           0.66, 6.18, 12.0, 0.66, size=20, minimum=18, color=PURPLE, family="Play")

    # 09 — Other hiring paths.
    s = d.slide("Вакансии и короткие задания", "Дополнительные пути взаимодействия внутри той же платформы.",
                "Вакансии имеют обязательную вилку RUB и статусы отклика. Pending-отклик можно отозвать. "
                "Задание назначается по специализации раз в неделю, кандидат присылает решение или подход. "
                "Ответ доступен автору компании без контактов, затем оценивается 1–5. Оценка задания не меняет грейд.",
                section="ПРОДУКТ")
    d.card(s, 0.55, 2.23, 5.97, 4.49, "КЛАССИЧЕСКИЙ НАЙМ", "Вакансия → отклик",
           "Компания публикует условия и вилку.\n\nСпециалист отправляет письмо.\n\nСтатусы: отправлен → просмотрен → принят / отказ.")
    d.card(s, 6.81, 2.23, 5.97, 4.49, "РЕГУЛЯРНАЯ АКТИВНОСТЬ", "Задание → обратная связь",
           "Одна назначенная задача на неделю.\n\nРешение или подход к решению.\n\nОценка 1–5 и комментарий компании.", fill=PALE)

    # 10 — Auth screenshot and integrations.
    s = d.slide("Вход и интеграции", "Локальная регистрация и серверный OAuth/OIDC с PKCE.",
                "Слева настоящий скриншот публичной страницы входа. Email подтверждается одноразовым письмом. "
                "ФСП ID и VK ID подключаются credentials сервера; explicit linking не объединяет аккаунты по email. "
                "Keycloak — локальный/демо OIDC провайдер. Реестр достижений сейчас синтетический. "
                "Внешние OAuth-токены остаются на сервере.", section="ИНТЕГРАЦИИ")
    d.rect(s, 0.55, 2.14, 4.6, 4.72)
    d.image(s, ASSETS / "login.png", 0.72, 2.2, 4.26, 4.59, crop=(490, 15, 1110, 885))
    items = [("Email / пароль", "Подтверждение почты и серверная сессия"),
             ("ФСП ID / VK ID", "Authorization Code + PKCE; явная привязка"),
             ("Реестр ФСП", "Отдельный импорт демонстрационных достижений"),
             ("PDF / SMTP", "Импорт резюме и письма подтверждения")]
    for i, (title, body) in enumerate(items):
        y = 2.18 + i * 1.16
        d.rect(s, 5.43, y, 7.35, 1.01, WHITE if i % 2 == 0 else PALE)
        d.text(s, title, 5.7, y + 0.13, 6.81, 0.41, size=22, minimum=20,
               family="Play", bold=True, color=PURPLE)
        d.text(s, body, 5.7, y + 0.61, 6.81, 0.27, size=15.5, minimum=14.5, color=MUTED)

    # 11 — Architecture.
    s = d.slide("Архитектура: модульный монолит", "Разделённые доменные контексты и общая транзакционная PostgreSQL.",
                "SvelteKit работает через Ktor API. Шесть контекстов разделены на domain, application, "
                "infrastructure и http; Ktor DI связывает порты с адаптерами. Общая БД сохраняет атомарность "
                "при записи результата и отзыве согласия. ArchitectureTest проверяет границы.", section="ТЕХНОЛОГИИ")
    d.flow(s, [("SvelteKit / Bun", "UI и API-клиент"), ("Kotlin / Ktor", "Use cases и доменные политики"),
               ("PostgreSQL", "Exposed · HikariCP · Flyway")], 2.23)
    contexts = [("identity", "Сессии и OAuth"), ("profiles", "Профили и согласия"),
                ("assessment", "CAT / IRT"), ("matching", "Рейтинг и объяснения"),
                ("recruitment", "Предложения и задачи"), ("achievements", "Реестр ФСП")]
    for i, (name, label) in enumerate(contexts):
        x, y = 0.55 + (i % 3) * 4.17, 4.05 + (i // 3) * 1.31
        d.rect(s, x, y, 3.89, 1.06, PALE)
        d.text(s, name, x + 0.22, y + 0.13, 3.45, 0.37,
               size=24, family="Play", bold=True, color=PURPLE)
        d.text(s, label, x + 0.22, y + 0.65, 3.45, 0.29,
               size=15.5, minimum=14, color=MUTED)

    # 12 — Data model.
    s = d.slide("Данные и бизнес-инварианты", "Результаты оценки связаны с профилем; взаимодействия — с обеими сторонами.",
                "Это укрупнённая схема, не полный ERD. Accounts связан с identities и role-profile. "
                "assessment_items принадлежат attempt. Вакансии, отклики, задания и приглашения "
                "защищены FK и проверками владельца. Грейд профиля выводится из истории текущей специализации. "
                "Полная модель и миграции находятся в техническом PDF.", section="ТЕХНОЛОГИИ")
    nodes = [(0.55, 2.25, "accounts + identities", "Роль и способы входа"),
             (0.55, 3.73, "candidate_profiles", "Опыт, контакты, согласия"),
             (0.55, 5.21, "assessment_attempts", "Результаты + items"),
             (5.13, 2.25, "employer_companies", "Компания и потребность"),
             (5.13, 3.73, "invitations / applications", "Связь с кандидатом"),
             (5.13, 5.21, "vacancies / short_tasks", "Предложения и задания")]
    for x, y, title, body in nodes:
        d.rect(s, x, y, 4.23, 1.12)
        d.text(s, title, x + 0.2, y + 0.16, 3.83, 0.4, size=22, minimum=18,
               family="Play", bold=True, color=PURPLE)
        d.text(s, body, x + 0.2, y + 0.73, 3.83, 0.26, size=14.5, minimum=13, color=MUTED)
    d.line(s, 2.66, 3.37, 2.66, 3.73)
    d.line(s, 2.66, 4.85, 2.66, 5.21)
    d.line(s, 4.78, 2.81, 5.13, 2.81)
    d.line(s, 4.78, 4.29, 5.13, 4.29)
    d.line(s, 9.36, 2.81, 9.5, 2.81)
    d.line(s, 9.5, 2.81, 9.5, 5.77)
    d.line(s, 9.5, 4.29, 9.36, 4.29)
    d.line(s, 9.5, 5.77, 9.36, 5.77)
    d.rect(s, 9.65, 2.25, 3.13, 4.39, PALE)
    d.text(s, "Контроль состояния", 9.88, 2.55, 2.67, 0.85,
           size=24, minimum=19.5, family="Play", bold=True, color=PURPLE)
    d.text(s, "Одна открытая попытка\n\nОтклик без дублей\n\nЗадание на неделю\n\nРоль и владение",
           9.88, 3.74, 2.67, 2.5, size=18.5, minimum=17)

    # 13 — Delivery.
    s = d.slide("Доставка: от коммита до работающего релиза", "GitHub Actions → GHCR → GitOps → Flux.",
                "CI запускает web/API/manifest checks. Main публикует три образа и атомарно фиксирует digests. "
                "Flux применяет конфигурацию в namespace fsp-career. CI не хранит cluster credentials. "
                "Promotion в Git и реальный rollout — разные стадии, состояние проверяется отдельно.", section="ТЕХНОЛОГИИ")
    d.flow(s, [("main", "Коммит проекта"), ("CI", "Web + API + manifests"),
               ("GHCR", "Образы с digest"), ("GitOps", "Promotion-коммит"),
               ("Flux", "Kubernetes rollout")], 2.48)
    d.card(s, 0.55, 4.24, 5.97, 2.53, "ЛОКАЛЬНО", "Docker Compose",
           "Web, API, базы, Keycloak и Mailpit")
    d.card(s, 6.81, 4.24, 5.97, 2.53, "РАБОЧЕЕ ОКРУЖЕНИЕ", "K3s / Envoy Gateway",
           "HTTPS, namespace fsp-career и Flux", fill=PALE)

    # 14 — Verification.
    s = d.slide("Проверки встроены в процесс разработки", "Доменные правила, HTTP-сценарии и доставка проверяются отдельно.",
                "Расскажите про уровни проверок без выдуманных чисел покрытия. В CI PostgreSQL 18, "
                "есть HTTP tests регистрации, OAuth, оценки, matching, вакансий и задач. "
                "Архитектурные проверки защищают слои; deploy tests проверяют manifests и promotion. "
                "Без TEST_DATABASE_URL локальные DB integration tests пропускаются.", section="КАЧЕСТВО")
    d.card(s, 0.55, 2.25, 3.89, 3.75, "WEB", "Типы и интерфейс",
           "Svelte / TypeScript\n\nBun regression tests\n\nESLint и production build")
    d.card(s, 4.72, 2.25, 3.89, 3.75, "API", "Правила и данные",
           "Unit и ArchitectureTest\n\nHTTP + PostgreSQL\n\nOAuth с MockEngine")
    d.card(s, 8.89, 2.25, 3.89, 3.75, "GITOPS", "Релиз и manifests",
           "Тесты promotion\n\nKustomize validation\n\nFlux health checks")
    d.text(s, "Примеры регрессий: CSRF · пересдачи · приватность · контакты · статусы откликов",
           0.65, 6.3, 12.0, 0.5, size=19.5, minimum=18.5, color=PURPLE, family="Play")

    # 15 — Scope and roadmap.
    s = d.slide("Текущая версия и развитие", "Работающий прототип со сквозными пользовательскими сценариями.",
                "Чётко отделите реализованные механизмы от внешних подключений. Синтетический реестр ФСП "
                "и Mailpit являются текущими демо-адаптерами. Направления развития — реальные интеграции, "
                "внешняя почта, экспертная калибровка на реальных данных и наблюдаемость. "
                "Не заявляйте официальную верификацию соревнований или измеренные бизнес-метрики.", section="РАЗВИТИЕ")
    d.card(s, 0.55, 2.25, 5.97, 4.48, "РЕАЛИЗОВАНО", "Полный путь взаимодействия",
           "Профиль, оценка и категория\n\nПодбор, приглашение и контакт\n\nВакансии, отклики и задания\n\nДемо, CI/CD и документация")
    d.card(s, 6.81, 2.25, 5.97, 4.48, "СЛЕДУЮЩИЕ ШАГИ", "Расширение интеграций",
           "Реальный реестр ФСП\n\nВнешний SMTP\n\nЭкспертная калибровка оценки\n\nНаблюдаемость и эксплуатация", fill=PALE)

    # 16 — Links and closing.
    s = d.slide("", "", "Завершите главным результатом: навыки подтверждаются, подбор объясняется, "
                "контакты раскрываются по явному действию. Откройте demo по QR или кликабельной ссылке. "
                "В PPTX также доступны ссылки на репозиторий, PDF и оба хранилища скринкастов.",
                dark=True, section="ДЕМОНСТРАЦИЯ")
    d.text(s, "Посмотрите полный сценарий", 0.65, 1.1, 8.4, 1.23,
           size=43, minimum=38, family="Play", bold=True, color=WHITE)
    d.text(s, "Навыки подтверждаются.\nПодбор объясняется.\nКонтакт открывается по решению.",
           0.7, 2.76, 8.1, 1.65, size=29, minimum=27, color=ROSE, family="Play")
    links = [("Открыть демо", DEMO), ("Репозиторий проекта", REPO),
             ("Техническая документация PDF", f"{REPO}/blob/main/docs/FSP-Career-documentation.pdf"),
             ("Скринкасты: Яндекс Диск", YANDEX)]
    for i, (label, url) in enumerate(links):
        d.text(s, label, 0.72, 4.9 + i * 0.39, 4.1 if i == 3 else 8.0, 0.31,
               size=17, minimum=16, color=WHITE, url=url)
    d.text(s, "Google Диск", 5.2, 6.07, 2.0, 0.3, size=17, minimum=16,
           color=ROSE, url=GOOGLE)
    qr = qrcode.QRCode(box_size=12, border=2)
    qr.add_data(DEMO)
    qr.make(fit=True)
    stream = BytesIO()
    qr.make_image(fill_color=f"#{PURPLE}", back_color="#FFFFFF").save(stream, format="PNG")
    stream.seek(0)
    d.rect(s, 9.33, 2.36, 3.18, 3.18, WHITE)
    picture = s.shapes.add_picture(stream, Inches(9.53), Inches(2.56), Inches(2.78), Inches(2.78))
    picture.click_action.hyperlink.address = DEMO
    d.text(s, "brokenbrain.online", 9.13, 5.84, 3.57, 0.45,
           size=18, minimum=17, color=WHITE, align=PP_ALIGN.CENTER, url=DEMO)

    d.finish(output)


def write_fontconfig(output):
    config = ET.Element("fontconfig")
    include = ET.SubElement(config, "include", ignore_missing="yes")
    include.text = "/etc/fonts/fonts.conf"
    ET.SubElement(config, "dir").text = str(output.parent / "fonts")
    ET.SubElement(config, "cachedir").text = str(output.parent / "font-cache")
    ET.ElementTree(config).write(output.parent / "fontconfig.xml", encoding="utf-8", xml_declaration=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.output)
    write_fontconfig(args.output)
