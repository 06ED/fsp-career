"""Generate an editable, brand-aligned FSP Career prototype pitch deck."""
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'docs' / 'FSP-Career-pitch-v4.pptx'
VIOLET = RGBColor(64, 47, 255)
DARK = RGBColor(27, 28, 33)
LIGHT = RGBColor(237, 237, 237)
GRAY = RGBColor(142, 143, 150)
WHITE = RGBColor(255, 255, 255)
P = Presentation()
P.slide_width, P.slide_height = Inches(13.333), Inches(7.5)


def textbox(slide, x, y, w, h, text, size=22, color=DARK, bold=False, font='Arial', line=1.15):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for idx, item in enumerate(text.split('\n')):
        par = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        par.text = item
        par.font.name = font
        par.font.size = Pt(size)
        par.font.bold = bold
        par.font.color.rgb = color
        par.line_spacing = line
        par.space_after = Pt(5)
    return shape


def rect(slide, x, y, w, h, color, outline=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background() if outline is None else None
    if outline:
        shp.line.color.rgb = outline
    return shp


def circle(slide, x, y, r, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(r), Inches(r))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    return shp


def slide_base(index, dark=False, section='ПЛАТФОРМА / ФСП КАРЬЕРА'):
    s = P.slides.add_slide(P.slide_layouts[6])
    back = DARK if dark else LIGHT
    foreground = WHITE if dark else DARK
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = back
    rect(s, .55, .48, .035, 6.52, VIOLET)
    textbox(s, .78, .42, 10, .35, section, 11, VIOLET if not dark else LIGHT, True, 'JetBrains Mono')
    textbox(s, 12.05, .42, .8, .4, f'{index:02d} / 07', 10, VIOLET if not dark else LIGHT, False, 'JetBrains Mono')
    textbox(s, .8, 7.07, 11, .22, 'ФЕДЕРАЦИЯ СПОРТИВНОГО ПРОГРАММИРОВАНИЯ  ·  СИНТЕТИЧЕСКИЙ ПРОТОТИП', 9, GRAY, False, 'JetBrains Mono')
    return s, foreground


def bullets(slide, items, x=.85, y=2.5, w=11.7, size=22, gap=.86, color=DARK):
    for i, text in enumerate(items):
        circle(slide, x, y + i * gap + .17, .12, VIOLET)
        textbox(slide, x + .38, y + i * gap, w - .4, gap - .07, text, size, color)


s, fg = slide_base(1)
textbox(s, .84, 1.35, 11.7, 2.5, 'КАРЬЕРА,\nПОДТВЕРЖДЁННАЯ ДЕЛОМ.', 48, DARK, True)
rect(s, .86, 4.52, 7.8, .035, VIOLET)
textbox(s, .85, 4.86, 10.6, 1.15, 'Кандидат проходит оценку → получает категорию → работодатель видит обоснованный подбор → отправляет приглашение.', 24, VIOLET)
circle(s, 11.42, 5.15, .95, VIOLET)

s, fg = slide_base(2, section='01 / ПРОБЛЕМА')
textbox(s, .85, 1.27, 11.5, 1.12, 'Резюме говорит. Навыки — доказывают.', 40, DARK, True)
bullets(s, [
    'Самоописание в резюме не равно подтверждённой квалификации.',
    'Поиск по словам теряет сильных специалистов и перегружает команды.',
    'Контактные данные и условия предложения требуют доверия и контроля.'
], y=2.72, size=23, gap=1.13)

s, fg = slide_base(3, dark=True, section='02 / МЕХАНИКА')
textbox(s, .85, 1.2, 10.8, .7, 'ДВА ПУТИ. ОДИН ПОНЯТНЫЙ МЭТЧ.', 35, WHITE, True)
for i, (head, body) in enumerate([
    ('01  СПЕЦИАЛИСТ', 'Профиль, специализация и тест. Категория определяется результатом оценки, не заявленным грейдом.'),
    ('02  РАБОТОДАТЕЛЬ', 'Потребность команды, категория и ранжированный список с объяснением каждого совпадения.'),
    ('03  КОНТАКТ', 'Приглашение с диапазоном зарплаты в рублях. Контакты — только после принятия кандидатом.')
]):
    x = .85 + i * 4.05
    rect(s, x, 2.55, 3.7, 3.5, RGBColor(43, 43, 52))
    textbox(s, x + .24, 2.85, 3.22, .75, head, 17, RGBColor(169, 162, 255), True, 'JetBrains Mono')
    textbox(s, x + .24, 3.72, 3.2, 2.0, body, 19, WHITE)

s, fg = slide_base(4, section='03 / ПРАВИЛА ПРОДУКТА')
textbox(s, .85, 1.23, 11.7, .88, 'Проверяемое ядро прототипа', 40, DARK, True)
bullets(s, [
    'Грейд не понижается принудительно; повторная попытка ограничена 90 днями.',
    'Счёт матча: оценка + подтверждённые записи достижений (≤20) + совпадение отрасли (+5).',
    'Публикация добровольна; до принятия приглашения контакты скрыты на уровне API.',
    'ФСП ID и официальная проверка достижений пока не интегрированы.'
], y=2.35, size=20, gap=1.02)

s, fg = slide_base(5, dark=True, section='04 / АРХИТЕКТУРА')
textbox(s, .85, 1.2, 10.4, .8, 'Простая и прозрачная архитектура', 36, WHITE, True)
for i, (title, body) in enumerate([
    ('SVELTEKIT', 'TypeScript · StyleX · GSAP · Bun\nDocker · кабинет двух ролей'),
    ('KTOR', 'Kotlin · JWT · OpenAPI\nОценка, ранжирование, доступ'),
    ('ДАННЫЕ', 'PostgreSQL · Flyway\nKeycloak · локальный Mailpit')
]):
    x = .85 + i * 4.05
    rect(s, x, 2.52, 3.75, 3.16, VIOLET if i == 1 else RGBColor(43, 43, 52))
    textbox(s, x + .25, 2.82, 3.25, .6, title, 18, WHITE, True, 'JetBrains Mono')
    textbox(s, x + .25, 3.65, 3.2, 1.65, body, 18, WHITE)

s, fg = slide_base(6, section='05 / ВАЛИДАЦИЯ — ЧЕСТНЫЙ СТАТУС')
textbox(s, .85, 1.18, 11.5, .82, 'Прототип ≠ валидированный инструмент', 37, DARK, True)
bullets(s, [
    '6 Docker-сервисов; 8 unit- и 1 DB-тест прошли; синтетический E2E через Keycloak пройден.',
    'Текущий банк заданий ограничен: необходимы экспертная проверка и сопоставимость форм.',
    'До реального найма: дискриминативность, повторяемость грейда, top-k и оценка пар экспертами.',
    'Никаких заявлений о точности подбора без независимых измерений.'
], y=2.27, size=20, gap=1.03)

s, fg = slide_base(7, dark=True, section='06 / ДЕМОНСТРАЦИЯ И ДАЛЬШЕ')
textbox(s, .85, 1.2, 11.3, 1.0, 'От профиля до принятого оффера.', 39, WHITE, True)
textbox(s, .85, 2.5, 10.9, 1.5, '01 Профиль  →  02 Тест  →  03 Категория\n04 Подбор  →  05 Приглашение  →  06 Контакт', 23, WHITE, False, 'JetBrains Mono')
rect(s, .85, 4.58, 11.55, .035, VIOLET)
textbox(s, .85, 4.9, 11, 1.28, 'Следующий этап: калибровка заданий, измеренная валидность, официальный коннектор ФСП ID и защищённое развёртывание.', 22, RGBColor(169, 162, 255))

P.save(OUT)
print(f'Created {OUT} ({len(P.slides)} slides)')
